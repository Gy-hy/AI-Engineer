"""Extractor 测试——分发、资源管控、超时、并发、异常兜底。

fixture：extract 吃的是【路径】，所以需要磁盘文件——用 tmp_path 生成
（每个测试文件自包含，不 import 其他测试文件）。

注意：本文件依赖真实插件注册（extract 按扩展名分发），所以在文件顶部
import 真实插件，且【不加】clear fixture——注册是 import 的副作用
（收集阶段发生一次、之后只读），与 test_plugins.py 同理。
"""

import io
import time
from typing import BinaryIO

import docx
import pytest
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from doc_extractor import Document, DocumentReader
from doc_extractor.core.exceptions import (
    DocumentReaderError,
    ExtractError,
    ExtractTimeoutError,
    UnsupportedFormatError,
)
from doc_extractor.core.extractor import Extractor

# 导入真实插件 → 触发自动注册（本文件测试依赖它们）
from doc_extractor.plugins.docx_reader import DocxReader  # noqa: F401
from doc_extractor.plugins.excel_reader import ExcelReader  # noqa: F401
from doc_extractor.plugins.pdf_reader import PdfReader  # noqa: F401


# ---------- 磁盘 fixture 工厂 ----------

def _write(path, buf: io.BytesIO) -> None:
    path.write_bytes(buf.getvalue())


def _pdf_bytes(text: str) -> io.BytesIO:
    """一页含指定文本的 PDF（reportlab 默认字体不支持中文，文本用 ASCII）。"""
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 750, text)
    c.save()
    buf.seek(0)
    return buf


def _docx_bytes(text: str) -> io.BytesIO:
    buf = io.BytesIO()
    d = docx.Document()
    d.add_paragraph(text)
    d.save(buf)
    buf.seek(0)
    return buf


def _xlsx_bytes() -> io.BytesIO:
    buf = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.append(["name", "age"])
    ws.append(["alice", 30])
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------- 探针插件（只活在测试里，不进产品路径） ----------

class SlowReader(DocumentReader):
    """超时探针插件：睡 0.5 秒再返回。"""

    extensions = (".slow",)

    def read(self, stream: BinaryIO) -> Document:
        time.sleep(0.5)
        return Document(text="")


class GatedReader(DocumentReader):
    """并发探针插件：read 期间 active +1，peak 记录历史最大并发。

    用类属性而不是实例属性——插件被多次实例化，探针状态要全局共享。
    （真实插件必须无状态，测试探针是例外，且不进产品路径。）
    """

    extensions = (".gate",)
    active = 0
    peak = 0

    def read(self, stream: BinaryIO) -> Document:
        type(self).active += 1
        type(self).peak = max(type(self).peak, type(self).active)
        time.sleep(0.05)
        type(self).active -= 1
        return Document(text="")


CASES = [
    pytest.param(".pdf", lambda: _pdf_bytes("hello pdf"), "hello pdf", id="pdf"),
    pytest.param(".docx", lambda: _docx_bytes("hello docx"), "hello docx", id="docx"),
    pytest.param(".xlsx", lambda: _xlsx_bytes(), "alice", id="xlsx"),
]


class TestExtract:
    @pytest.mark.parametrize("ext, make_buf, expected", CASES)
    def test_extract(self, tmp_path, ext, make_buf, expected):
        """三种格式都能正确抽取（验收 2 的自动化部分）。"""
        p = tmp_path / f"sample{ext}"
        _write(p, make_buf())

        with Extractor() as ex:
            doc = ex.extract(p)

        assert expected in doc.text

    def test_unsupported_format_lists_supported(self, tmp_path):
        """验收 3：不支持的扩展名 → UnsupportedFormatError，
        消息同时包含该扩展名和已支持格式列表。"""
        p = tmp_path / "note.unknown"
        p.write_text("hi")

        with pytest.raises(UnsupportedFormatError) as exc:
            with Extractor() as ex:
                ex.extract(p)

        msg = str(exc.value)
        assert ".unknown" in msg          # 包含这个扩展名
        assert ".pdf" in msg              # 包含已支持格式列表

    def test_missing_file_cause_chain(self, tmp_path):
        """验收 7：非插件错误（文件不存在）也被兜底，cause 链完整。"""
        p = tmp_path / "missing.pdf"

        with pytest.raises(ExtractError) as exc:
            with Extractor() as ex:
                ex.extract(p)

        assert isinstance(exc.value.cause, FileNotFoundError)
        assert exc.value.__cause__ is exc.value.cause   # from e 建立的链

    def test_bad_file_wrapped(self, tmp_path):
        """验收 7：插件抛 DocumentReaderError → 调用方收到 ExtractError。"""
        p = tmp_path / "bad.pdf"
        p.write_bytes(b"not a real pdf")

        with pytest.raises(ExtractError) as exc:
            with Extractor() as ex:
                ex.extract(p)

        assert isinstance(exc.value.cause, DocumentReaderError)

    def test_timeout(self, tmp_path):
        """验收 5：超时 → ExtractTimeoutError，不无限等待。

        诚实细节：超时后线程池里的任务还会继续跑完（线程无法强杀），
        所以本测试结束时 SlowReader 可能仍在睡——close() 的 wait=True
        会等它跑完再释放，干净退出。
        """
        p = tmp_path / "slow.slow"
        p.write_bytes(b"x")

        ex = Extractor(timeout=0.05)
        with pytest.raises(ExtractTimeoutError):
            ex.extract(p)
        ex.close()

    def test_max_concurrency_one(self, tmp_path):
        """验收 6：max_concurrency=1 时，任何瞬间只有一个 read 在执行
        （探针插件的 peak 计数验证）。"""
        GatedReader.active = 0
        GatedReader.peak = 0
        paths = []
        for i in range(3):
            p = tmp_path / f"g{i}.gate"
            p.write_bytes(b"x")
            paths.append(p)

        ex = Extractor(max_concurrency=1)
        ex.extract_many(paths)
        ex.close()

        assert GatedReader.peak == 1

    def test_extract_many_order(self, tmp_path):
        """批量抽取：结果顺序与输入一致（8.5 决策 4 的正确实现路径）。"""
        paths = []
        for i in range(3):
            p = tmp_path / f"m{i}.pdf"
            _write(p, _pdf_bytes(f"hello pdf {i}"))
            paths.append(p)

        ex = Extractor()
        docs = ex.extract_many(paths)
        ex.close()

        for i, doc in enumerate(docs):
            assert f"hello pdf {i}" in doc.text