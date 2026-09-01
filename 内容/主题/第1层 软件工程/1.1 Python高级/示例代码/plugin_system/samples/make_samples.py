"""生成 samples/ 下的三个样例文件（CLI 演示用）。

运行：uv run python samples/make_samples.py

关于 PDF 中文：reportlab 默认字体 Helvetica 不含中文字形；用内置 CID
字体（STSong-Light）虽能显示中文，但 pypdf 抽取 CID 字体文本往往得不到
可检索的 Unicode（缺少 ToUnicode 映射）。样例 PDF 是给 CLI 演示抽取用的，
所以保持英文文本。DOCX/XLSX 对 Unicode 天然支持，用中文。
"""

from pathlib import Path

import docx
from openpyxl import Workbook
from reportlab.pdfgen import canvas

SAMPLES_DIR = Path(__file__).parent


def make_pdf() -> None:
    c = canvas.Canvas(str(SAMPLES_DIR / "sample.pdf"))
    c.setTitle("Doc Extractor Sample")
    c.drawString(100, 750, "Hello from sample.pdf")
    c.drawString(100, 720, "This text was extracted by the PDF plugin.")
    c.save()


def make_docx() -> None:
    d = docx.Document()
    d.add_heading("Doc Extractor 样例文档", level=1)
    d.add_paragraph("这是一个 DOCX 样例，用于验证插件抽取。")
    d.add_paragraph("包含段落和表格：")
    table = d.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "语言"
    table.rows[0].cells[1].text = "用途"
    table.rows[1].cells[0].text = "Python"
    table.rows[1].cells[1].text = "AI 系统工程师"
    d.save(str(SAMPLES_DIR / "sample.docx"))


def make_xlsx() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "成绩单"
    ws.append(["姓名", "分数", "备注"])
    ws.append(["张三", 95, "优秀"])
    ws.append(["李四", 88, '=IF(B3>=90,"优秀","良好")'])
    # 注意：上面是公式单元格。data_only=True 抽取时读【缓存值】，
    # 而缓存值由 Excel 软件保存时写入——openpyxl 生成的文件没有缓存值，
    # 所以这个单元格抽取出来会是空的。这正是 excel_reader.py 注释里说的坑。
    wb.save(str(SAMPLES_DIR / "sample.xlsx"))


if __name__ == "__main__":
    make_pdf()
    make_docx()
    make_xlsx()
    print("已生成:", [p.name for p in sorted(SAMPLES_DIR.glob("sample.*"))])