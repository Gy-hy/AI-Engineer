"""DOCX 插件——用 python-docx 解析 Word 文档为 Document。

职责边界与 pdf_reader 相同：只解析、不管资源、靠继承自动注册。
"""

from typing import BinaryIO

import docx

from ..core.document import Document
from ..core.exceptions import DocumentReaderError
from ..core.reader import DocumentReader


class DocxReader(DocumentReader):
    """DOCX → Document。"""

    extensions = (".docx",)

    def read(self, stream: BinaryIO) -> Document:
        """按 DocumentReader.read 的契约实现：不关 stream、无状态、失败抛 DocumentReaderError。"""
        try:
            doc = docx.Document(stream)
        except Exception as e:
            # 损坏文件 / 不是 zip 容器时，python-docx 抛 PackageNotFoundError 等——
            # 同样包装成插件作者的语言，保留 cause 链。
            raise DocumentReaderError(f"DOCX 解析失败: {e}") from e

        parts: list[str] = []

        # 段落
        for p in doc.paragraphs:
            if p.text.strip():          # 跳过空行（Word 文档里常见的分隔空段落）
                parts.append(p.text)

        # 表格：表格里的文字【不在】doc.paragraphs 里，必须单独取。
        # 行内单元格用制表符连接，保持"一行"的结构感。
        # 已知取舍：这里先取完所有段落再取所有表格，
        # 段落与表格在原文中的交错顺序会丢失（真实保序需要遍历 body 的
        # XML 子元素按出现顺序处理——留作扩展练习）。
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                parts.append("\t".join(cells))

        return Document(text="\n".join(parts))